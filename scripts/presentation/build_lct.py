# -*- coding: utf-8 -*-
"""Презентация Kostik в шаблоне ЛЦТ 2026 (template/lct_template.pptx).

Слайды 7–11 шаблона — обязательный блок: сетку и оформление не трогаем, только текст и картинки в существующих
полях; на слайде карточек команды лишние карточки удаляются целиком. Дальше — описательная часть решения на фонах,
плашках и в палитре шаблона (Montserrat), скриншоты — в рамке браузера. Остальные слайды шаблона удаляются.

    /tmp/pv/bin/python build_lct.py  →  Kostik_LCT2026.pptx
"""
import copy
import os

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
A = lambda f: os.path.join(HERE, 'assets', f)
TODO = '⟨уточнить⟩'

PINK, PINK_L, LAV, PURPLE, DEEP, INK, WHITE, GREY = (RGBColor.from_string(x) for x in
    ('FF0053', 'FFD6E3', '8A83D1', '520977', '2D1451', '1C1D22', 'FFFFFF', '6B6E78'))
FONT = 'Montserrat'

prs = Presentation(os.path.join(HERE, 'template', 'lct_template.pptx'))
S = list(prs.slides)
LIGHT_SRC, PURPLE_SRC = S[9], S[8]          # фон «светлый» (слайд 10) и «фиолетовый» (слайд 9)
PILL_SRC = [sh for sh in S[9].shapes if sh.name == 'Скругленный прямоугольник 1'][0]


# ------------------------------------------------------------------ помощники
def put(shape, lines, size=None, bold=None, color=None):
    """Заменить текст фигуры, сохранив оформление первого абзаца и первого фрагмента."""
    tf = shape.text_frame
    p0 = tf.paragraphs[0]
    ppr = copy.deepcopy(p0._p.pPr) if p0._p.pPr is not None else None
    rpr = None
    for p in tf.paragraphs:
        for r in p.runs:
            rpr = copy.deepcopy(r._r.rPr) if r._r.rPr is not None else None
            break
        if rpr is not None:
            break
    for p in list(tf.paragraphs)[1:]:
        p._p.getparent().remove(p._p)
    for r in list(p0.runs):
        r._r.getparent().remove(r._r)
    for el in list(p0._p):
        if el.tag in (qn('a:br'), qn('a:fld')):
            p0._p.remove(el)
    for i, line in enumerate(lines if isinstance(lines, list) else [lines]):
        p = p0 if i == 0 else tf.add_paragraph()
        if i and ppr is not None:
            p._p.insert(0, copy.deepcopy(ppr))
        segs = line if isinstance(line, list) else [(line, {})]
        for text, st in segs:
            r = p.add_run()
            if rpr is not None:
                r._r.insert(0, copy.deepcopy(rpr))
            r.text = text
            if size or st.get('size'):
                r.font.size = Pt(st.get('size', size))
            if bold is not None or 'bold' in st:
                r.font.bold = st.get('bold', bold)
            if color is not None or 'color' in st:
                r.font.color.rgb = st.get('color', color)


def place(ph, f, mode='cover'):
    """Вставить картинку в поле шаблона, сохранив его место и размер (python-pptx сбрасывает их на позицию макета)."""
    from PIL import Image as _I
    L, T, W, H = ph.left, ph.top, ph.width, ph.height
    pic_ = ph.insert_picture(A(f))
    iw, ih = _I.open(A(f)).size
    ri, rb = iw / ih, W / H
    pic_.crop_left = pic_.crop_right = pic_.crop_top = pic_.crop_bottom = 0
    if mode == 'cover':
        pic_.left, pic_.top, pic_.width, pic_.height = L, T, W, H
        if ri > rb:
            e = (1 - rb / ri) / 2; pic_.crop_left = pic_.crop_right = e
        else:
            e = (1 - ri / rb) / 2; pic_.crop_top = pic_.crop_bottom = e
    else:
        if ri > rb:
            w, h = W, int(W / ri)
        else:
            w, h = int(H * ri), H
        pic_.left, pic_.top, pic_.width, pic_.height = L, T + (H - h) // 2, w, h
    return pic_


def drop(shape):
    shape._element.getparent().remove(shape._element)


def by_name(slide, name):
    return [sh for sh in slide.shapes if sh.name == name]


def new_slide(dark=False, title=''):
    src = PURPLE_SRC if dark else LIGHT_SRC
    s = prs.slides.add_slide(src.slide_layout)
    for ph in list(s.placeholders):
        if ph.placeholder_format.type not in (1,):          # оставить только заголовок
            drop(ph)
    bg = src._element.cSld.bg
    if bg is not None:
        nb = copy.deepcopy(bg)
        for blip in nb.iter(qn('a:blip')):
            part = src.part.related_part(blip.get(qn('r:embed')))
            blip.set(qn('r:embed'), s.part.relate_to(part, RT.IMAGE))
        s._element.cSld.insert(0, nb)
    pill = copy.deepcopy(PILL_SRC._element)
    s.shapes._spTree.insert(2, pill)
    width = Inches(max(3.2, 0.7 + 0.215 * len(title)))
    pill.spPr.xfrm.ext.cx = width
    tp = [ph for ph in s.placeholders if ph.placeholder_format.type == 1]
    if tp:
        t = tp[0]
        t.left, t.top, t.width, t.height = Inches(0.56), Inches(0.49), width - Inches(0.3), Inches(0.41)
        t.text_frame.text = title
        for r in t.text_frame.paragraphs[0].runs:
            r.font.color.rgb = WHITE
    return s


def text(s, x, y, w, h, lines, size=14, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font=FONT, spacing=None):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, m, 0)
    for i, line in enumerate(lines if isinstance(lines, list) else [lines]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.space_after = Pt(spacing)
        for seg, st in (line if isinstance(line, list) else [(line, {})]):
            r = p.add_run()
            r.text = seg
            f = r.font
            f.name = st.get('font', font)
            f.size = Pt(st.get('size', size))
            f.bold = st.get('bold', bold)
            f.color.rgb = st.get('color', color)
    return tb


def box(s, x, y, w, h, fill=WHITE, line=None, radius=0.08):
    sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    sh.adjustments[0] = radius
    sh.fill.solid()
    sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(1.25)
    sh.shadow.inherit = False
    return sh


def pic(s, f, x, y, w=None, h=None):
    return s.shapes.add_picture(A(f), Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)


def badge(s, n, x, y, color=PINK, d=0.46):
    c = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(d), Inches(d))
    c.fill.solid(); c.fill.fore_color.rgb = color; c.line.fill.background(); c.shadow.inherit = False
    tf = c.text_frame
    for m in ('margin_left', 'margin_right', 'margin_top', 'margin_bottom'):
        setattr(tf, m, 0)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = str(n); r.font.name = FONT; r.font.size = Pt(14); r.font.bold = True; r.font.color.rgb = WHITE


def notes(s, t):
    s.notes_slide.notes_text_frame.text = t


# ------------------------------------------------------------------ 7. Титул
s7 = S[6]
for sh in list(s7.placeholders):
    t = sh.placeholder_format.type
    if t == 3:
        put(sh, 'Квантовый Скачок')
    elif t == 2:
        put(sh, ['Сервис искусственного интеллекта по оценке качества исследований плотности костей человека · решение Kostik'])
    elif t == 18:
        place(sh, 'logo_task.png', 'contain')

# ------------------------------------------------------------------ 8. Краткое описание и команда
s8 = S[7]
for ph in [ph for ph in s8.placeholders if ph.placeholder_format.type == 18]:
    place(ph, 'team_pair.png')
for sh in list(s8.shapes):
    if sh.is_placeholder and sh.has_text_frame and sh.placeholder_format.type == 18:
        continue
    elif sh.is_placeholder and sh.has_text_frame and sh.placeholder_format.type == 1:
        put(sh, 'Kostik', color=PURPLE)
    elif sh.has_text_frame and sh.text_frame.text.startswith('В чем суть'):
        put(sh, ['Проверяет снимки денситометрии (DXA) поясничного отдела и бедра сразу после исследования: область, '
                 'вердикт 0/1, тип нарушения по ТЗ и объяснение прямо на снимке. Работает локально, без интернета.'])
    elif sh.has_text_frame and sh.text_frame.text.startswith('Что делает'):
        put(sh, ['Объяснимость для врача: атлас с разметкой, граф «измерение → норма → вердикт», честная пометка сомнений; '
                 'устойчивость к любым файлам и работа офлайн — на сервере и на ноутбуке.'])
    elif sh.has_text_frame and sh.text_frame.text.startswith('Капитан'):
        tf = sh.text_frame
        vals = {'Капитан:': f' Юрий Коноплёв, {TODO}', 'Кол-во участников:': ' 2 человека', 'Краткое описание:': '',
                'как образовалась команда?': TODO, 'место работы/учебы участников?': TODO, 'Город и регион:': f' {TODO}'}
        for p in tf.paragraphs:
            full = ''.join(r.text for r in p.runs).strip()
            key = next((k for k in vals if full.startswith(k)), None)
            if key is None:
                continue
            if key in ('как образовалась команда?', 'место работы/учебы участников?'):
                p.runs[0].text = vals[key]
                for r in p.runs[1:]:
                    r.text = ''
            else:
                label = p.runs[0]
                label.text = key
                for r in p.runs[1:]:
                    r.text = ''
                if vals[key]:
                    extra = p.runs[-1] if len(p.runs) > 1 else p.add_run()
                    extra.text = vals[key]

# ------------------------------------------------------------------ 9. Карточки команды: лишние три карточки — удалить целиком
s9 = S[8]
names = {'Скругленный прямоугольник 58', 'Скругленный прямоугольник 61', 'Скругленный прямоугольник 64'}
tb_x = {Emu(Inches(5.61)).real, Emu(Inches(8.13)).real, Emu(Inches(10.65)).real}
for sh in list(s9.shapes):
    if sh.left is None:
        if sh.is_placeholder and sh.has_text_frame and sh.placeholder_format.type == 1:
            put(sh, 'КОМАНДА')
        continue
    x = round(Emu(sh.left).inches, 2) if sh.left is not None else None
    if sh.name in names or (sh.shape_type == 17 and x in (5.61, 8.13, 10.65)) or \
       (sh.is_placeholder and sh.placeholder_format.type == 18 and x in (5.68, 8.22, 10.72)):
        drop(sh)
team = [dict(x=0.57, pic_x=0.64, name='Юрий Коноплёв', role='Капитан, руководитель проекта, продукт', nick='@bimodaling', photo='team_yuri_sq.png'),
        dict(x=3.09, pic_x=3.16, name='Алексей Чуркин', role='Стенд, Telegram-бот, проверка сервиса', nick='@lesha_cfc', photo='team_alexey_sq.png')]
pics = [(round(Emu(ph.left).inches, 2), ph) for ph in s9.placeholders if ph.placeholder_format.type == 18 and ph.left is not None]
for x, ph in pics:
    for m in team:
        if x == m['pic_x']:
            place(ph, m['photo'])
for sh in list(s9.shapes):
    if sh.left is None:
        if sh.is_placeholder and sh.has_text_frame and sh.placeholder_format.type == 1:
            put(sh, 'КОМАНДА')
        continue
    x = round(Emu(sh.left).inches, 2)
    for m in team:
        if sh.has_text_frame and x == m['x'] and sh.text_frame.text.startswith('Имя'):
            put(sh, m['name'])
        elif sh.has_text_frame and x == m['x'] and sh.text_frame.text.startswith('Роль'):
            put(sh, [m['role'], 'Telegram ' + m['nick'], 'Телефон: ' + TODO, TODO + ' (работа/учёба)'])
    if sh.is_placeholder and sh.has_text_frame and sh.placeholder_format.type == 1:
        put(sh, 'КОМАНДА')

# ------------------------------------------------------------------ 10. История команды
s10 = S[9]
for sh in s10.shapes:
    t = sh.text_frame.text if sh.has_text_frame else ''
    if sh.is_placeholder and sh.placeholder_format.type == 1:
        put(sh, 'О КОМАНДЕ')
    elif t.startswith('Расскажите, как вы собрались'):
        put(sh, [f'{TODO}: как собрались и где работали вместе. Решение сделали вдвоём за две недели; код, тесты '
                 'и документацию писали вместе с ИИ-агентом Claude через Codellake.'])
    elif t.startswith('Что вас вдохновило'):
        put(sh, ['Брак укладки DXA ведёт к ошибкам в диагнозе остеопороза, а критерии качества в ТЗ формализуемы — '
                 'можно сделать объяснимый сервис, которому поверит врач.'])
    elif t.startswith('Расскажите о самых интересных'):
        put(sh, ['В DICOM нет масштаба и области — определяем их по изображению. Из 499 файлов уникальны 252 — боролись с утечкой данных. '
                 'Брака бедра всего 41 — честно показываем предел модели. Рёбра похожи на предмет — сервис помечает сомнение.'])

# ------------------------------------------------------------------ 11. Коротко о решении
s11 = S[10]
for sh in s11.placeholders:
    t = sh.text_frame.text
    if t.startswith('Опишите в чем'):
        put(sh, ['Поясница — правила по анатомическим ориентирам: угол оси, охват, контраст пятна вне столба.',
                 'Бедро — ExtraTrees по 47 признакам формы, инференс на numpy.',
                 'Контейнер без сети, API, MCP, приложение для Windows и Linux.',
                 '0,31 с на снимок на 2 ядрах CPU; 100 % файлов, отказ — с причиной.'])
    elif t.startswith('Опишите ваши идеи'):
        put(sh, ['Для кого: рентгенолаборанты и врачи отделений денситометрии.',
                 'Эффект: брак виден, пока пациент в кабинете, — меньше повторных исследований.',
                 'Пилот в 1–2 отделениях, дообучение на правках врачей из пульта.',
                 'Далее — интеграция с PACS/ЕРИС и отчёт о качестве отделений.'])

# ------------------------------------------------------------------ описательная часть
# 12. Проблема
s = new_slide(title='ПРОБЛЕМА')
text(s, 0.56, 1.3, 12.2, 0.6, 'Каждый третий снимок поясницы — с браком укладки', size=26, bold=True, color=DEEP)
stats = [('32 %', 'снимков поясничного отдела — с нарушением по оценке экспертов организатора', PINK),
         ('27 %', 'снимков бедра — с нарушением укладки или полей зоны интереса', PURPLE),
         ('0', 'автоматических проверок: качество смотрят вручную и выборочно', LAV)]
for i, (b, t, c) in enumerate(stats):
    x = 0.56 + i * 4.15
    box(s, x, 2.25, 3.9, 2.85, line=LAV)
    text(s, x + 0.35, 2.45, 3.3, 1.1, b, size=54, bold=True, color=c)
    text(s, x + 0.35, 3.65, 3.25, 1.3, t, size=14, color=INK)
text(s, 0.56, 5.5, 12.2, 1.0, [[('Цена ошибки: ', {'bold': True, 'color': DEEP}),
      ('наклон оси, неполный охват или металл в кадре искажают плотность кости — остеопороз пропускают или лечат лишнее, пациента вызывают повторно.', {})]], size=15)
notes(s, 'В обучающем наборе эксперты отметили нарушения у 32 из 99 поясниц и 41 из 150 бёдер.')

# 13. Решение — дашборд в браузере
s = new_slide(title='РЕШЕНИЕ')
text(s, 0.56, 1.3, 12.2, 0.6, 'Секунда после исследования — вердикт по каждому снимку', size=24, bold=True, color=DEEP)
pic(s, 'browser/dash.png', 0.3, 1.95, w=7.9)
pts = [('Загрузка', 'DICOM, папка или zip — перетащить в рамку'), ('Дашборд', 'сколько годных, где брак, время на снимок'),
       ('Галерея', 'красная рамка — нарушение, причина под снимком'), ('Выгрузка', 'таблица по п. 2.5 ТЗ: XLSX, CSV, ZIP разметки')]
for i, (h, t) in enumerate(pts):
    y = 2.25 + i * 1.1
    badge(s, i + 1, 8.5, y)
    text(s, 9.1, y - 0.02, 3.9, 1.0, [[(h, {'bold': True, 'size': 16, 'color': DEEP})], [(t, {'size': 13})]])
notes(s, 'Дашборд проверки: синтетический фантом, 6 снимков, один с посторонним предметом.')

# 14. Объяснимость — карточка в браузере
s = new_slide(dark=True, title='ОБЪЯСНИМОСТЬ')
text(s, 0.56, 1.3, 12.2, 0.6, 'Врач видит причину прямо на снимке', size=24, bold=True, color=WHITE)
pic(s, 'browser/card.png', 0.3, 1.95, w=7.9)
pts = [('Атлас', 'позвонки Th12–L5, ось, гребни, предмет в рамке'), ('Атлас ⇄ снимок', 'одно нажатие — чистый рентген без разметки'),
       ('Граф решения', 'измерение → норма → вердикт; или текстом'), ('Сомнение', '«предмет? или ребро — проверьте»')]
for i, (h, t) in enumerate(pts):
    y = 2.25 + i * 1.1
    badge(s, i + 1, 8.5, y, PINK)
    text(s, 9.1, y - 0.02, 3.9, 1.0, [[(h, {'bold': True, 'size': 16, 'color': WHITE})], [(t, {'size': 13, 'color': PINK_L})]])

# 15. Архитектура
s = new_slide(title='АРХИТЕКТУРА')
text(s, 0.56, 1.3, 12.2, 0.6, 'Один конвейер для стенда, API, контейнера и приложения', size=24, bold=True, color=DEEP)
steps = [('Вход', 'DICOM, папка или zip; сжатые форматы; дубли по пикселям'), ('Область', 'поясница, левое или правое бедро'),
         ('Проверки', 'поясница — правила по ориентирам; бедро — ExtraTrees'), ('Пояснение', 'атлас на снимке, граф решения, текст и голос'),
         ('Выход', 'CSV / XLSX по ТЗ 2.5, ZIP разметки, API, MCP')]
for i, (h, t) in enumerate(steps):
    x = 0.56 + i * 2.5
    box(s, x, 2.2, 2.25, 2.6, fill=PINK_L if i == 2 else WHITE, line=LAV)
    badge(s, i + 1, x + 0.22, 2.4, PINK if i == 2 else PURPLE)
    text(s, x + 0.22, 3.0, 1.9, 0.4, h, size=17, bold=True, color=DEEP)
    text(s, x + 0.22, 3.45, 1.85, 1.3, t, size=12)
    if i < 4:
        text(s, x + 2.25, 3.2, 0.25, 0.4, '→', size=18, color=LAV, align=PP_ALIGN.CENTER)
stack = [('Стенд', 'https://ltz2026.ru/'), ('Контейнер по ТЗ', 'build.sh · run.sh · serve.sh, без сети'),
         ('Приложение 1.0-Beta', 'Windows, Ubuntu, Debian — офлайн'), ('Интеграции', 'HTTP API · MCP · Telegram-бот')]
for i, (h, t) in enumerate(stack):
    text(s, 0.56 + i * 3.1, 5.35, 2.95, 0.9, [[(h, {'bold': True, 'color': DEEP})], [(t, {'color': GREY})]], size=13)

# 16. Данные
s = new_slide(title='ДАННЫЕ')
text(s, 0.56, 1.3, 12.2, 0.6, 'Честная оценка — без утечки данных', size=24, bold=True, color=DEEP)
st = [('100', 'исследований'), ('252', 'уникальных снимка'), ('247', 'дублей убрано'), ('99 / 150', 'поясниц / бёдер с оценкой')]
for i, (b, t) in enumerate(st):
    x = 0.56 + i * 3.1
    text(s, x, 2.05, 3.0, 0.9, b, size=44, bold=True, color=PINK)
    text(s, x, 2.95, 3.0, 0.4, t, size=14, color=INK)
rows = [('Без утечки', 'фолды — по исследованиям; общие снимки нескольких исследований — всегда в одной группе'),
        ('Честная оценка', '5 фолдов × 50 повторов; пороги подбираются только на обучающей части'),
        ('Дисбаланс', 'брак 32 % и 27 %: веса классов, порог по F1, метрики F1 и ROC-AUC'),
        ('Интервалы', '95 % — по повторам и бутстрепом по исследованиям')]
for i, (h, t) in enumerate(rows):
    y = 3.75 + i * 0.62
    text(s, 0.56, y, 2.6, 0.5, h, size=15, bold=True, color=DEEP)
    text(s, 3.2, y, 9.6, 0.5, t, size=14)

# 17. Результаты — диаграмма
s = new_slide(dark=True, title='РЕЗУЛЬТАТЫ')
text(s, 0.56, 1.3, 12.2, 0.6, 'Позвоночник — уверенно, бедро — подсказка для проверки', size=24, bold=True, color=WHITE)
labels = ['Охват', 'Посторонние предметы', 'Поля ROI (бедро)', 'Поясница, итог', 'Наклон оси', 'Бедро, итог', 'Укладка бедра']
vals = [0.88, 0.85, 0.80, 0.75, 0.75, 0.66, 0.62]
cd = CategoryChartData(); cd.categories = list(reversed(labels)); cd.add_series('ROC-AUC', list(reversed(vals)))
box(s, 0.45, 1.95, 7.9, 5.0, fill=WHITE)
gf = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, Inches(0.6), Inches(2.05), Inches(7.6), Inches(4.8), cd)
ch = gf.chart
ch.has_legend = False
ch.has_title = True
ch.chart_title.text_frame.text = 'ROC-AUC · кросс-валидация по исследованиям'
tr = ch.chart_title.text_frame.paragraphs[0].runs[0]; tr.font.size = Pt(13); tr.font.name = FONT; tr.font.color.rgb = GREY
pl = ch.plots[0]; pl.gap_width = 45; pl.has_data_labels = True
dl = pl.data_labels; dl.number_format = '0.00'; dl.number_format_is_linked = False; dl.position = XL_LABEL_POSITION.INSIDE_END
dl.font.size = Pt(13); dl.font.bold = True; dl.font.color.rgb = WHITE; dl.font.name = FONT
va = ch.value_axis; va.minimum_scale = 0.5; va.maximum_scale = 1.0; va.major_unit = 0.1
va.tick_labels.font.size = Pt(11); va.tick_labels.font.color.rgb = GREY; va.major_gridlines.format.line.color.rgb = RGBColor(0xE5, 0xE7, 0xE9)
ca = ch.category_axis; ca.tick_labels.font.size = Pt(13); ca.tick_labels.font.name = FONT; ca.tick_labels.font.color.rgb = INK
ca.format.line.color.rgb = RGBColor(0xE5, 0xE7, 0xE9)
for i, v in enumerate(reversed(vals)):
    pt = pl.series[0].points[i]; pt.format.fill.solid()
    pt.format.fill.fore_color.rgb = PURPLE if v >= 0.8 else LAV if v >= 0.7 else PINK
legend = [(PURPLE, '≥ 0,80', 'предметы и охват — можно доверять'), (LAV, '0,70–0,80', 'итог поясницы и ось: сколиоз похож на наклон'),
          (PINK, '< 0,70', 'бедро — мало размеченного брака, проверить глазами')]
for i, (c, h, t) in enumerate(legend):
    y = 2.05 + i * 1.62
    box(s, 8.6, y, 4.35, 1.42, fill=WHITE)
    b = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(8.85), Inches(y + 0.22), Inches(0.3), Inches(0.3))
    b.fill.solid(); b.fill.fore_color.rgb = c; b.line.fill.background(); b.shadow.inherit = False
    text(s, 9.3, y + 0.15, 3.5, 0.45, h, size=18, bold=True, color=DEEP)
    text(s, 8.85, y + 0.65, 3.95, 0.7, t, size=12.5)
notes(s, 'Охват 0,88, предметы 0,85, поля ROI 0,80 (7 примеров), поясница 0,75, ось 0,75, бедро 0,66, укладка 0,62.')

# 18. Метрики — таблица
s = new_slide(title='МЕТРИКИ ПО П. 8.4 ТЗ')
text(s, 0.56, 1.3, 12.2, 0.5, 'Кросс-валидация: 5 фолдов по исследованиям × 50 повторов; порог — только на обучающей части', size=14, color=GREY)
rows = [['Критерий', 'Брак / всего', 'Чувствит.', 'Специф.', 'Сбаланс. точн.', 'F1 (95 % ДИ)', 'ROC-AUC (95 % ДИ)'],
        ['Поясничный отдел, итог', '32 / 99', '0,69', '0,67', '0,68', '0,58 (0,46–0,74)', '0,75 (0,64–0,84)'],
        ['   охват', '6 / 99', '0,67', '0,98', '0,82', '0,69 (0,40–0,93)', '0,88 (0,58–1,00)'],
        ['   наклон оси', '10 / 99', '0,41', '0,86', '0,63', '0,30 (0,11–0,60)', '0,75 (0,55–0,91)'],
        ['   посторонние предметы', '17 / 99', '0,68', '0,84', '0,76', '0,55 (0,34–0,72)', '0,85 (0,76–0,94)'],
        ['Бедро, итог', '41 / 150', '0,65', '0,57', '0,61', '0,47', '0,66 (0,62–0,69)'],
        ['   укладка и ротация', '36 / 150', '—', '—', '—', '—', '0,62'],
        ['   поля зоны интереса', '7 / 150', '—', '—', '—', '—', '0,80']]
tb = s.shapes.add_table(len(rows), 7, Inches(0.56), Inches(1.95), Inches(12.2), Inches(3.9)).table
widths = [3.0, 1.4, 1.2, 1.2, 1.5, 1.95, 1.95]
for j, w in enumerate(widths):
    tb.columns[j].width = Inches(w)
for i, r in enumerate(rows):
    for j, v in enumerate(r):
        c = tb.cell(i, j); c.text = v
        p = c.text_frame.paragraphs[0]; p.alignment = PP_ALIGN.LEFT if j == 0 else PP_ALIGN.CENTER
        f = p.runs[0].font; f.name = FONT; f.size = Pt(12 if i else 11.5)
        c.fill.solid()
        if i == 0:
            c.fill.fore_color.rgb = PURPLE; f.bold = True; f.color.rgb = WHITE
        else:
            c.fill.fore_color.rgb = WHITE if i % 2 else RGBColor(0xF6, 0xF2, 0xFA)
            f.color.rgb = INK; f.bold = r[0] in ('Поясничный отдел, итог', 'Бедро, итог')
            if j == 6:
                v0 = float(v.split()[0].replace(',', '.'))
                f.bold = True; f.color.rgb = PURPLE if v0 >= 0.8 else LAV if v0 >= 0.7 else PINK
box(s, 0.56, 6.1, 12.2, 0.72, fill=PINK_L)
text(s, 0.85, 6.1, 11.7, 0.72, [[('Сводная фитнес-функция 0,58  ', {'bold': True, 'color': DEEP}),
      ('= 0,35·F1 + 0,35·ROC-AUC по областям + 0,30·macro-F1 типов; базовая линия 0,570 (0,477–0,644)', {})]], size=13, anchor=MSO_ANCHOR.MIDDLE)

# 19. Эксперименты
s = new_slide(title='ЭКСПЕРИМЕНТЫ')
text(s, 0.56, 1.3, 12.2, 0.6, 'Главная победа — детектор предметов: ROC-AUC 0,55 → 0,85', size=24, bold=True, color=DEEP)
rows = [['Модель для бедра', 'ROC-AUC'], ['Логистическая регрессия, 21 признак', '0,60'], ['Признаки ResNet-50 / DINOv2', 'до 0,64'],
        ['Признаки RAD-DINO', '0,64'], ['Бустинг · каскадный лес · Viola–Jones', 'не выше 0,66'],
        ['ExtraTrees, 47 признаков — выбрана', '0,66'], ['ExtraTrees + RAD-DINO', '0,67 — в пределах ДИ']]
tb = s.shapes.add_table(len(rows), 2, Inches(0.56), Inches(2.1), Inches(6.6), Inches(3.9)).table
tb.columns[0].width = Inches(4.5); tb.columns[1].width = Inches(2.1)
for i, r in enumerate(rows):
    for j, v in enumerate(r):
        c = tb.cell(i, j); c.text = v; f = c.text_frame.paragraphs[0].runs[0].font; f.name = FONT; f.size = Pt(13)
        c.fill.solid(); c.fill.fore_color.rgb = PURPLE if i == 0 else (PINK_L if 'выбрана' in r[0] else WHITE)
        f.color.rgb = WHITE if i == 0 else INK; f.bold = i == 0 or 'выбрана' in r[0]
lessons = [(PURPLE, 'Предметы: 0,55 → 0,85', 'локальный контраст пятна вместо порога яркости'),
           (LAV, 'Ось: сколиоз', 'изгиб столба похож на наклон; три другие меры оси — хуже'),
           (PINK, 'Бедро: мало сигнала', 'нужна разметка ротации по малому вертелу и масштаб в мм')]
for i, (c, h, t) in enumerate(lessons):
    y = 2.1 + i * 1.35
    box(s, 7.5, y, 5.3, 1.18, fill=WHITE, line=c)
    text(s, 7.8, y + 0.14, 4.8, 0.45, h, size=17, bold=True, color=c if c != LAV else PURPLE)
    text(s, 7.8, y + 0.6, 4.8, 0.5, t, size=13)

# 20. Скорость
s = new_slide(dark=True, title='СКОРОСТЬ')
text(s, 0.56, 1.3, 12.2, 0.6, 'Секунда на исследование при допуске ТЗ в 3 минуты', size=24, bold=True, color=WHITE)
st = [('0,31 с', 'на снимок · 2 ядра CPU, без видеокарты'), ('90 с', 'весь обучающий набор, 252 снимка'),
      ('100 %', 'файлов обработано, отказ — с причиной'), ('191 МБ', 'пик памяти процесса')]
for i, (b, t) in enumerate(st):
    x = 0.56 + i * 3.1
    box(s, x, 2.15, 2.9, 2.2, fill=WHITE)
    text(s, x + 0.25, 2.35, 2.5, 0.9, b, size=40, bold=True, color=PINK if i == 0 else DEEP)
    text(s, x + 0.25, 3.3, 2.45, 0.95, t, size=13)
req = [('Минимум', '2 ядра, 4 ГБ памяти, 64-битный Linux; образ контейнера 2,3 ГБ'),
       ('Сеть', 'не нужна: контейнер запускается с --network none'),
       ('Воспроизводимость', 'версии зафиксированы, повторный прогон совпадает побайтно')]
for i, (h, t) in enumerate(req):
    y = 4.8 + i * 0.58
    text(s, 0.56, y, 3.0, 0.5, h, size=15, bold=True, color=WHITE)
    text(s, 3.6, y, 9.2, 0.5, t, size=14, color=PINK_L)

# 21. Если файл не тот
s = new_slide(title='УСТОЙЧИВОСТЬ')
text(s, 0.56, 1.3, 12.2, 0.6, 'Любой файл — понятный ответ, без падений', size=24, bold=True, color=DEEP)
pic(s, 'browser/reject.png', 0.3, 1.95, w=7.9)
pts = [('175 чужих файлов', 'КТ, МРТ, УЗИ, цветные, маски, битые zip — 0 падений'),
       ('Распознавание', '«КТ позвоночника, вид сбоку», надписи и водяные знаки'),
       ('Отчёт денситометра', 'снимок вырезается и проверяется с пометкой'),
       ('Данные организатора', '252 снимка — результат прежний, 0 отказов')]
for i, (h, t) in enumerate(pts):
    y = 2.25 + i * 1.1
    badge(s, i + 1, 8.5, y, PURPLE)
    text(s, 9.1, y - 0.02, 3.9, 1.0, [[(h, {'bold': True, 'size': 16, 'color': DEEP})], [(t, {'size': 13})]])

# 22. Продукт
s = new_slide(title='ПРОДУКТ')
text(s, 0.56, 1.3, 12.2, 0.6, 'Готово к использованию уже сейчас', size=24, bold=True, color=DEEP)
pic(s, 'browser/download.png', 0.3, 1.95, w=7.9)
pts = [('Веб-стенд', 'https://ltz2026.ru/ — проверка и примеры'), ('Приложение', '.exe · .msi · .deb — работает офлайн'),
       ('Контейнер', 'сборка и запуск одной командой, как в ТЗ'), ('Интеграции', 'HTTP API · MCP · Telegram-бот')]
for i, (h, t) in enumerate(pts):
    y = 2.25 + i * 1.1
    badge(s, '✓', 8.5, y, PURPLE)
    text(s, 9.1, y - 0.02, 3.9, 1.0, [[(h, {'bold': True, 'size': 16, 'color': DEEP})], [(t, {'size': 13})]])

# 23. План развития
s = new_slide(title='ПЛАНЫ ПО РАЗВИТИЮ')
text(s, 0.56, 1.3, 12.2, 0.6, 'Честно о пределах — и как их закрыть в пилоте', size=24, bold=True, color=DEEP)
lim = ['Бедро — подсказка: ROC-AUC 0,66', 'Сколиоз путается с наклоном оси', 'Th12 определяется приблизительно', 'Поля в сантиметрах: в DICOM нет масштаба']
box(s, 0.56, 2.1, 4.4, 4.6, fill=PINK_L)
text(s, 0.85, 2.3, 3.9, 0.5, 'Ограничения', size=18, bold=True, color=PINK)
text(s, 0.85, 2.95, 3.9, 3.6, [[('• ' + t, {})] for t in lim], size=14, spacing=10)
plan = [('Пилот', '1–2 отделения: проверка каждого снимка в день исследования'), ('Разметка', 'врачи правят спорные вердикты в пульте — данные для дообучения'),
        ('Бедро и Th12', 'разметка малого вертела и уровней, масштаб от аппарата'), ('Масштаб', 'интеграция с PACS / ЕРИС, отчёт о качестве отделений')]
for i, (h, t) in enumerate(plan):
    y = 2.1 + i * 1.18
    badge(s, i + 1, 5.4, y, PINK)
    text(s, 6.05, y - 0.02, 6.8, 1.05, [[(h, {'bold': True, 'size': 17, 'color': DEEP})], [(t, {'size': 14})]])

# 24. Финал
s = new_slide(dark=True, title='ПОПРОБУЙТЕ САМИ')
text(s, 0.56, 1.5, 7.5, 1.0, 'Kostik', size=60, bold=True, color=WHITE)
text(s, 0.56, 2.6, 7.3, 1.3, 'Контроль качества снимков денситометрии — за доли секунды, прямо на снимке и без интернета.', size=20, color=PINK_L)
text(s, 0.56, 4.3, 7.3, 1.4, [[('Команда «Квантовый Скачок»', {'bold': True, 'color': WHITE, 'size': 18})],
                               [('Юрий Коноплёв · Алексей Чуркин', {'color': PINK_L})],
                               [('Telegram: @bimodaling · @lesha_cfc', {'color': LAV})]], size=15)
box(s, 8.9, 1.55, 3.9, 4.55, fill=WHITE)
pic(s, 'qr.png', 9.25, 1.8, w=3.2)
link = text(s, 8.9, 5.2, 3.9, 0.6, 'https://ltz2026.ru/', size=18, bold=True, color=PURPLE, align=PP_ALIGN.CENTER)
link.text_frame.paragraphs[0].runs[0].hyperlink.address = 'https://ltz2026.ru/'

# ------------------------------------------------------------------ порядок: 7–11 обязательные, затем наши; прочие — удалить
sld = prs.slides._sldIdLst
ids = list(sld)
keep_first = ids[6:11]
ours = ids[37:]
for el in ids:
    sld.remove(el)
for el in keep_first + ours:
    sld.append(el)
for el in ids:
    if el not in keep_first and el not in ours:
        prs.part.drop_rel(el.get(qn('r:id')))
out = os.path.join(HERE, 'Kostik_LCT2026.pptx')
prs.save(out)
print('готово:', out, len(prs.slides._sldIdLst), 'слайдов')
